#!/usr/bin/env python3
"""Build the Furniture Pickup Go/No-Go workbook."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = (
    Path(__file__).resolve().parents[1]
    / "products"
    / "furniture-pickup-sheet"
    / "build"
    / "Furniture-Pickup-Go-No-Go.xlsx"
)

INK = "192722"
GREEN = "175B3F"
INPUT = "D9EAD3"
OUTFILL = "FFF2CC"
MUTED = "58665E"
GO_FILL = "C6EFCE"
GO_INK = "176B37"
NOGO_FILL = "F4C7C3"
NOGO_INK = "9C1F1F"
thin = Border(
    left=Side(style="thin", color="D5DDD5"),
    right=Side(style="thin", color="D5DDD5"),
    top=Side(style="thin", color="D5DDD5"),
    bottom=Side(style="thin", color="D5DDD5"),
)
title_font = Font(name="Calibri", size=18, bold=True, color=INK)
label_font = Font(name="Calibri", size=11, color=INK)
head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
note_font = Font(name="Calibri", size=10, italic=True, color=MUTED)
money = '"$"#,##0.00'
pct = "0%"
num = "0.00"


def style_input(cell, number=None):
    cell.fill = PatternFill("solid", fgColor=INPUT)
    cell.font = label_font
    cell.border = thin
    cell.alignment = Alignment(vertical="center")
    if number:
        cell.number_format = number


def style_out(cell, number=None):
    cell.fill = PatternFill("solid", fgColor=OUTFILL)
    cell.font = Font(name="Calibri", size=11, bold=True, color=INK)
    cell.border = thin
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if number:
        cell.number_format = number


def col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put_label(ws, row, text):
    cell = ws.cell(row, 1, text)
    cell.font = label_font
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28
    return cell


def put_note(ws, row, text):
    cell = ws.cell(row, 3, text)
    cell.font = note_font
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    return cell


def section_header(ws, row, text, fill):
    cell = ws.cell(row, 1, text)
    cell.font = head_font
    cell.fill = PatternFill("solid", fgColor=fill)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    return cell


def add_pickup_sheet(wb):
    ws = wb.active
    ws.title = "Pickup"
    ws.sheet_properties.tabColor = GREEN
    ws["A1"] = "Furniture Pickup Go/No-Go"
    ws["A1"].font = title_font
    ws.merge_cells("A1:C1")
    ws["A2"] = (
        "Enter the green cells. Yellow cells calculate. Expected resale is YOUR GUESS. "
        "This sheet does not know local demand or what the piece will actually sell for. "
        "No sale or income is guaranteed. Starting numbers are a fictional Harbor Mill "
        "example. Overwrite them with THIS pickup."
    )
    ws["A2"].font = note_font
    ws.merge_cells("A2:C2")
    ws.row_dimensions[2].height = 48
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")

    for col, text in enumerate(["INPUTS - green cells only", "Value", "Note"], 1):
        cell = ws.cell(4, col, text)
        cell.font = head_font
        cell.fill = PatternFill("solid", fgColor=GREEN)

    inputs = [
        (5, "Piece / listing name", "Harbor Mill oak sideboard (fictional)", None,
         "Overwrite with this pickup. Starting values are the fictional example."),
        (6, "Where it is", "Marketplace", None,
         "Curb / Marketplace / Other. A listing photo is not the same as seeing the piece."),
        (7, "Pickup date", "12 Sep 2026", None,
         "Your date. This is not a calendar or a booking tool."),
        (8, "Asking price", 40, money,
         "What you would pay today. Type 0 for a free curb piece."),
        (9, "Cash on hand for this pickup", 80, money,
         "Cash you can actually spend today. Not next week's sale."),
        (10, "Estimated repair / parts", 25, money,
         "Glue, knobs, finish, pads. Type 0 only if you will sell it as-is."),
        (11, "Estimated hours", 4, num,
         "Drive both ways, load, repair, photos, listing, and the buyer meetup."),
        (12, "Your target hourly", 25, money,
         "What you need per hour for this work. The sheet does not pick it for you."),
        (13, "Round-trip miles", 12, num,
         "Home to piece to destination and back, as you will actually drive it."),
        (14, "Fuel per mile", 0.20, money,
         "What this trip costs you per mile. Not a published reimbursement rate."),
        (15, "Platform fees", 0.10, pct,
         "Percent you expect to pay the resale platform. Type 0% for a local cash sale."),
        (16, "Expected resale (YOUR GUESS)", 150, money,
         "A guess you typed. Not an appraisal. Not local demand. Not a sold price."),
    ]
    for row, label, value, fmt, note in inputs:
        put_label(ws, row, label)
        ws.cell(row, 2, value)
        style_input(ws.cell(row, 2), fmt)
        put_note(ws, row, note)

    section_header(ws, 18, "RED FLAGS - Yes or No", GREEN)
    flags = [
        (19, "Smoke odor", "No", "Smell the piece, not the listing. Yes forces NO-GO."),
        (20, "Bedbugs / pests", "No", "Look at joints and undersides. Yes forces NO-GO."),
        (21, "Structural damage (frame, rot, split)", "No",
         "A broken frame is not a hardware run. Yes forces NO-GO."),
        (22, "Missing hardware", "Yes",
         "Caps the decision at CAUTION. Set No only if replacement parts are already in repair $."),
        (23, "Too big for the vehicle you will use today", "No",
         "If it will not fit the vehicle that will show up, Yes forces NO-GO."),
    ]
    for row, label, value, note in flags:
        put_label(ws, row, label)
        ws.cell(row, 2, value)
        style_input(ws.cell(row, 2))
        put_note(ws, row, note)

    section_header(ws, 25, "THRESHOLDS - green cells, you may edit", GREEN)
    put_label(ws, 26, "GO if implied hourly is at least")
    ws["B26"] = "=B12"
    style_input(ws["B26"], money)
    put_note(ws, 26, "Default follows your target hourly. Type over the formula to set a different floor.")
    put_label(ws, 27, "CAUTION if implied hourly is at least")
    ws["B27"] = "=B12*0.5"
    style_input(ws["B27"], money)
    put_note(ws, 27, "Below this floor the money side is NO-GO, even with no red flags.")

    section_header(ws, 29, "RESULTS - do not type here", "7A5C12")

    results = [
        (30, "Fuel cost", "=B13*B14", money, "Miles x fuel per mile."),
        (31, "Platform fee on the guessed resale", "=B16*B15", money,
         "Guessed resale x the fee percent you typed."),
        (32, "Total cash out", "=B8+B10+B30+B31", money,
         "Ask + repair + fuel + fee. Time is not cash out; it shows up in implied hourly."),
        (33, "Money left after costs", "=B16-B32", money,
         "Guessed resale minus cash out. Still depends on a guess."),
        (34, "Implied hourly", '=IF(B11<=0,"",B33/B11)', money,
         "Money left / hours. If you under-count hours, this number lies."),
        (35, "Cash shortfall", "=MAX(0,B8-B9)", money,
         "How much you are short of the asking price today."),
        (36, "Hard red flags", '=(B19="Yes")+(B20="Yes")+(B21="Yes")+(B23="Yes")', "0",
         "Count of Yes on smoke, pests, structure, or too big. Missing hardware is separate."),
    ]
    for row, label, formula, fmt, note in results:
        put_label(ws, row, label)
        ws.cell(row, 2, formula)
        style_out(ws.cell(row, 2), fmt)
        put_note(ws, row, note)

    put_label(ws, 37, "DECISION")
    ws["B37"] = (
        '=IF(B11<=0,"NO-GO",IF(B8>B9,"NO-GO",IF(B33<=0,"NO-GO",'
        'IF(B36>0,"NO-GO",IF(B34<B27,"NO-GO",'
        'IF(OR(B34<B26,B22="Yes"),"CAUTION","GO"))))))'
    )
    style_out(ws["B37"])
    ws["B37"].font = Font(name="Calibri", size=13, bold=True, color=INK)
    ws["A37"].font = Font(name="Calibri", size=13, bold=True, color=INK)
    put_note(ws, 37, "GO / CAUTION / NO-GO from the thresholds and flags above. You can edit those.")

    ws.conditional_formatting.add(
        "B37",
        CellIsRule(
            operator="equal",
            formula=['"GO"'],
            fill=PatternFill("solid", fgColor=GO_FILL),
            font=Font(name="Calibri", size=13, bold=True, color=GO_INK),
        ),
    )
    ws.conditional_formatting.add(
        "B37",
        CellIsRule(
            operator="equal",
            formula=['"CAUTION"'],
            fill=PatternFill("solid", fgColor=OUTFILL),
            font=Font(name="Calibri", size=13, bold=True, color="7A5C12"),
        ),
    )
    ws.conditional_formatting.add(
        "B37",
        CellIsRule(
            operator="equal",
            formula=['"NO-GO"'],
            fill=PatternFill("solid", fgColor=NOGO_FILL),
            font=Font(name="Calibri", size=13, bold=True, color=NOGO_INK),
        ),
    )

    put_label(ws, 38, "Why")
    ws["B38"] = (
        '=IF(B11<=0,"Type estimated hours (drive, pickup, repair, photos, listing, meetup) before you trust a decision.",'
        'IF(B8>B9,"Asking is more than cash on hand for this pickup.",'
        'IF(B33<=0,"Your resale guess does not cover ask + repair + fuel + fees. That guess is still a guess.",'
        'IF(B36>0,"A hard red flag is Yes: smoke, pests, structural damage, or too big for today\'s vehicle.",'
        'IF(B34<B27,"Implied hourly is below your CAUTION threshold.",'
        'IF(AND(B34<B26,B22="Yes"),"Implied hourly is below your GO threshold, and hardware is missing.",'
        'IF(B34<B26,"Implied hourly is below your GO threshold.",'
        'IF(B22="Yes","Hourly meets your GO threshold, but missing hardware caps this at CAUTION.",'
        '"Implied hourly meets your GO threshold, cash covers the ask, and no red flags are Yes. Resale is still YOUR guess."))))))))'
    )
    style_out(ws["B38"])
    ws.merge_cells("B38:C38")
    ws.row_dimensions[38].height = 48

    ws["A40"] = (
        "Before you treat GO as a plan: photograph all sides, the underside, and the joints; "
        "smell the piece; measure the piece against the vehicle opening and the doorway it must pass. "
        "This workbook cannot see the piece, local demand, or the buyer who has not appeared yet."
    )
    ws["A40"].font = note_font
    ws.merge_cells("A40:C40")
    ws["A40"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[40].height = 48

    source_dv = DataValidation(
        type="list",
        formula1='"Curb,Marketplace,Other"',
        allow_blank=False,
    )
    source_dv.error = "Pick Curb, Marketplace, or Other."
    source_dv.errorTitle = "Where it is"
    ws.add_data_validation(source_dv)
    source_dv.add("B6")

    flag_dv = DataValidation(
        type="list",
        formula1='"Yes,No"',
        allow_blank=False,
    )
    flag_dv.error = "Pick Yes or No."
    flag_dv.errorTitle = "Red flag"
    ws.add_data_validation(flag_dv)
    flag_dv.add("B19:B23")

    ws["B16"].comment = Comment(
        "Type a resale guess you could live with being wrong. The sheet will not check sold comps or local demand.",
        "Bloom Shelf",
    )
    ws["B26"].comment = Comment(
        "Type over this cell if you want a GO floor different from your target hourly.",
        "Bloom Shelf",
    )

    col_widths(ws, [52, 22, 78])
    ws.freeze_panes = "A5"
    ws.print_title_rows = "1:2"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.oddFooter.left.text = (
        "Bloom Web Services LLC  not a demand forecast  support@thaliabloom.com"
    )


def add_example_sheet(wb):
    ws = wb.create_sheet("Worked example")
    ws["A1"] = "Fictional worked example - Harbor Mill oak sideboard"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")
    ws["A2"] = (
        "Made-up piece on a made-up street in a made-up town. Not a customer, "
        "not a sold listing, and not a recommended bid. The Pickup sheet ships "
        "with these same starter numbers so you can see the yellow cells move."
    )
    ws["A2"].font = note_font
    ws.merge_cells("A2:B2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 48

    ws["A4"] = "What they typed"
    ws["B4"] = "Value"
    ws["A4"].font = head_font
    ws["B4"].font = head_font
    ws["A4"].fill = PatternFill("solid", fgColor=GREEN)
    ws["B4"].fill = PatternFill("solid", fgColor=GREEN)

    example = [
        ("Piece", "Harbor Mill oak sideboard (fictional)"),
        ("Where", "Marketplace"),
        ("Asking price", "$40.00"),
        ("Cash on hand", "$80.00"),
        ("Estimated repair / parts", "$25.00 (two knobs + wood filler)"),
        ("Estimated hours", "4.00 (drive, load, repair, photos, list, meetup)"),
        ("Target hourly", "$25.00 (the owner typed this; it is not advice)"),
        ("Round-trip miles", "12.00"),
        ("Fuel per mile", "$0.20"),
        ("Platform fees", "10%"),
        ("Expected resale (YOUR GUESS)", "$150.00"),
        ("Smoke / pests / structure / too big", "No / No / No / No"),
        ("Missing hardware", "Yes"),
        ("GO threshold", "$25.00 (follows target hourly)"),
        ("CAUTION threshold", "$12.50 (half of target hourly)"),
    ]
    for i, (key, value) in enumerate(example, 5):
        ws.cell(i, 1, key).font = label_font
        ws.cell(i, 2, value).font = label_font

    ws["A21"] = "What the Pickup sheet should show with those inputs"
    ws["A21"].font = Font(name="Calibri", size=12, bold=True)
    ws.merge_cells("A21:B21")

    calc = [
        ("Fuel cost", "12.00 x $0.20 = $2.40"),
        ("Platform fee", "$150.00 x 10% = $15.00"),
        ("Total cash out", "$40.00 + $25.00 + $2.40 + $15.00 = $82.40"),
        ("Money left after costs", "$150.00 - $82.40 = $67.60"),
        ("Implied hourly", "$67.60 / 4.00 = $16.90"),
        ("Cash shortfall", "$0.00 (cash on hand covers the ask)"),
        ("Hard red flags", "0"),
        ("DECISION", "CAUTION"),
        ("Why", "Implied hourly is below the GO threshold, and hardware is missing."),
    ]
    for i, (key, value) in enumerate(calc, 22):
        ws.cell(i, 1, key).font = label_font
        ws.cell(i, 2, value).font = Font(name="Calibri", size=11, bold=True)

    ws["A32"] = (
        "A CAUTION is not a hidden GO. The $150 resale is a guess someone typed. "
        "If your Pickup sheet is more than a few cents off, check that platform fees "
        "is 10% not 10, that red flags are the dropdown values Yes and No, and that "
        "you did not type in a yellow cell."
    )
    ws["A32"].font = note_font
    ws.merge_cells("A32:B32")
    ws["A32"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[32].height = 48
    col_widths(ws, [52, 78])


def add_limits_sheet(wb):
    ws = wb.create_sheet("Limits")
    ws["A1"] = "Limits - read before you drive"
    ws["A1"].font = title_font
    lines = [
        "This workbook is a calculator for ONE pickup. It is not an appraisal, a demand report, a listing tool, a hauling service, or legal, tax, pest, or safety advice.",
        "Expected resale is whatever you type. The sheet does not know local demand, sold comps, season, or what a buyer will actually pay. A GO is not a sale.",
        "It does not inspect the piece. Photograph all sides, the underside, and the joints. Smell it. Measure the piece against the vehicle opening and the doorway it must pass.",
        "Smoke, bedbugs/pests, structural damage, and too-big-for-today's-vehicle are hard flags. Any one of them makes the decision NO-GO. Missing hardware caps the decision at CAUTION unless you set that flag to No after pricing the parts in repair $.",
        "GO and CAUTION hourly floors live in green cells on Pickup. Change them. The starter rule is GO at your target hourly and CAUTION at half of that. Below the CAUTION floor, the money side is NO-GO.",
        "Cash on hand is cash you can spend today. If the asking price is higher, the decision is NO-GO even when the implied hourly looks fine.",
        "Hours must include unpaid time: drive both ways, loading, repair, photos, listing, and the meetup. If you leave those out, implied hourly is a vanity number.",
        "No income, sale, or 'this piece will sell' promise. If the piece changes when you arrive, change the sheet or walk.",
        "Seller: Bloom Web Services LLC. Support: support@thaliabloom.com. 14-day refund by email.",
    ]
    for i, line in enumerate(lines, 3):
        ws.cell(i, 1, line).font = label_font
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 48
    ws.column_dimensions["A"].width = 110


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    add_pickup_sheet(wb)
    add_example_sheet(wb)
    add_limits_sheet(wb)
    names = wb.sheetnames
    if names != ["Pickup", "Worked example", "Limits"]:
        raise SystemExit(f"unexpected sheets: {names}")
    wb.save(OUT)
    print(OUT, OUT.stat().st_size, names)


if __name__ == "__main__":
    main()
